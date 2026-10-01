"""SPEC-V2 §7 — .xlsx import (server-side, openpyxl).

Proves the real deliverable: a genuine .xlsx file, built with openpyxl (not
a renamed CSV), round-trips through POST /api/import/analyze ->
POST /api/import/commit exactly like a CSV does — same dedupe-on-import
(D13), same explicit-consent step, same two-phase wizard. Also proves the
multi-sheet interrogation (D11: ask before parsing a single row) and that a
corrupt upload returns a structured 422, never a 500.

Runs standalone, same dev-auth harness as the other suites.
"""
from __future__ import annotations

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["EXCEEDBOX_DB"] = os.path.join(tempfile.mkdtemp(), "test_xlsx_import.db")
os.environ["EXCEEDBOX_UPLOAD_ROOT"] = tempfile.mkdtemp()
os.environ["EXCEEDBOX_DEV_AUTH"] = "1"
os.environ.pop("SUPABASE_JWT_SECRET", None)
os.environ["EXCEEDBOX_ALLOWED_EMAIL_DOMAINS"] = "test.example,contract.example,fable-fix.example,exceed-re.ae"

import openpyxl                       # noqa: E402
from fastapi.testclient import TestClient   # noqa: E402

from app import db, devauth           # noqa: E402
from app.api import app               # noqa: E402
from app import csvimport             # noqa: E402

client = TestClient(app)

EMAILS = {
    "admin": "xlsx-admin@test.example",
    "marketing": "xlsx-marketing@test.example",
    "sales": "xlsx-sales@test.example",
}


def hdr(role: str) -> dict:
    return {"Authorization": "Bearer " + devauth.mint(sub="xlsx-test-sub-%s" % role, email=EMAILS[role])}


def fresh():
    db.reset()
    con = db.connect()
    for role, email in EMAILS.items():
        con.execute("""INSERT INTO app_user (id,email,display_name,role,is_active)
                       VALUES (?,?,?,?,1)""", ("xlsx-%s" % role, email, role.title(), role))
    con.commit()
    con.close()


def _xlsx_bytes(headers, rows, sheet_title="Sheet1") -> bytes:
    """Builds a real .xlsx workbook in memory with openpyxl — never a
    hand-written byte string pretending to be one."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _multi_sheet_xlsx_bytes(sheets: dict) -> bytes:
    """`sheets` is {sheet_name: (headers, rows)}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── real .xlsx round-trip ─────────────────────────────────────────────────

def test_xlsx_round_trip_analyze_and_commit():
    fresh()
    data = _xlsx_bytes(
        ["Name", "Email", "Company"],
        [
            ["山田 太郎", "taro@xlsx-test.example", "Taro Excel Co"],
            ["Jane Smith", "jane@xlsx-test.example", "Smith Ventures"],
        ],
    )
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("leads.xlsx", data,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body.get("needs_sheet_selection") is not True
    assert body["rows_seen"] == 2
    assert set(body["columns"]) == {"Name", "Email", "Company"}
    assert body["guessed_mapping"]["Name"] == "name"
    assert body["guessed_mapping"]["Email"] == "email"
    assert body["preview_new"] == 2
    assert body["preview_duplicate"] == 0
    token = body["token"]

    r2 = client.post("/api/import/commit", headers=hdr("marketing"),
                     json={"token": token, "mapping": body["guessed_mapping"],
                           "answers": {"consent_basis": "explicit"}})
    assert r2.status_code == 200, r2.text
    committed = r2.json()
    assert committed["rows_created"] == 2
    assert committed["rows_skipped"] == 0
    assert committed["consent_basis_applied"] == "explicit"
    assert len(committed["lead_ids"]) == 2

    # dedupe on import (D13) — re-importing the SAME workbook must merge,
    # not duplicate, both rows.
    r3 = client.post("/api/import/analyze", headers=hdr("marketing"),
                     files={"file": ("leads-again.xlsx", data,
                                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r3.status_code == 200
    body3 = r3.json()
    assert body3["preview_new"] == 0
    assert body3["preview_duplicate"] == 2
    r4 = client.post("/api/import/commit", headers=hdr("marketing"),
                     json={"token": body3["token"], "mapping": body3["guessed_mapping"],
                           "answers": {"consent_basis": "unknown"}})
    assert r4.status_code == 200
    assert r4.json()["rows_created"] == 0
    assert r4.json()["rows_merged"] == 2


def test_xlsx_blank_rows_and_cells_are_skipped_not_crashed_on():
    fresh()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Email"])
    ws.append(["Real Person", "real@xlsx-test.example"])
    ws.append([None, None])            # a genuinely blank row
    ws.append(["", ""])                # blank strings, still effectively empty
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("blanks.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    # openpyxl's None-row is filtered; the ""/"" row still counts as a row
    # (real, if entirely blank, cells) — commit() is what skips no-name rows.
    assert r.json()["rows_seen"] >= 1


# ── multi-sheet workbook: interrogate BEFORE parsing a row (D11) ──────────

def test_xlsx_multi_sheet_asks_which_sheet_first():
    fresh()
    data = _multi_sheet_xlsx_bytes({
        "Dubai Leads": (["Name", "Email"], [["Dubai Person", "dubai@xlsx-test.example"]]),
        "Lombok Leads": (["Name", "Email"], [["Lombok Person", "lombok@xlsx-test.example"]]),
    })
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("two-sheets.xlsx", data,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["needs_sheet_selection"] is True
    assert set(body["sheets"]) == {"Dubai Leads", "Lombok Leads"}
    assert "token" not in body, "nothing may be stashed/committed-to before a sheet is chosen"

    # answering the interrogation: re-upload the SAME file with sheet_name set
    r2 = client.post("/api/import/analyze", headers=hdr("marketing"),
                     files={"file": ("two-sheets.xlsx", data,
                                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                     data={"sheet_name": "Lombok Leads"})
    assert r2.status_code == 200, r2.text
    body2 = r2.json()
    assert body2.get("needs_sheet_selection") is not True
    assert body2["rows_seen"] == 1
    assert body2["sample_rows"][0]["Name"] == "Lombok Person"

    r3 = client.post("/api/import/commit", headers=hdr("marketing"),
                     json={"token": body2["token"], "mapping": body2["guessed_mapping"],
                           "answers": {"consent_basis": "unknown"}})
    assert r3.status_code == 200, r3.text
    assert r3.json()["rows_created"] == 1


def test_xlsx_single_sheet_workbook_never_asks():
    fresh()
    data = _xlsx_bytes(["Name"], [["Solo Sheet Person"]], sheet_title="Only Sheet")
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("one-sheet.xlsx", data,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json().get("needs_sheet_selection") is not True
    assert r.json()["rows_seen"] == 1


# ── corrupt file: structured 422, never a 500 ─────────────────────────────

def test_corrupt_xlsx_returns_422_not_500():
    fresh()
    garbage = b"this is not a zip file, not an xlsx workbook, not anything openpyxl can open"
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("broken.xlsx", garbage,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 422, r.text
    body = r.json()
    assert body["error"]["code"] == "unparseable_file"
    assert "excel" in body["error"]["message"].lower() or "workbook" in body["error"]["message"].lower()


def test_xlsx_with_no_header_row_is_422_not_500():
    fresh()
    wb = openpyxl.Workbook()
    ws = wb.active
    buf = io.BytesIO()
    wb.save(buf)  # a real workbook, but genuinely empty — zero rows
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("empty.xlsx", buf.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 422, r.text
    assert r.json()["error"]["code"] == "unparseable_file"


def test_empty_upload_is_422_not_500():
    fresh()
    r = client.post("/api/import/analyze", headers=hdr("marketing"),
                    files={"file": ("nothing.xlsx", b"",
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 422, r.text
    assert r.json()["error"]["code"] == "empty_file"


# ── permissions: sales has no import.csv capability, same as CSV ─────────

def test_xlsx_import_denied_for_sales():
    fresh()
    data = _xlsx_bytes(["Name"], [["Nobody"]])
    r = client.post("/api/import/analyze", headers=hdr("sales"),
                    files={"file": ("x.xlsx", data,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 403


# ── direct unit coverage of csvimport.parse_xlsx_bytes() ──────────────────

def test_parse_xlsx_bytes_unit():
    fresh()
    data = _xlsx_bytes(["A", "B"], [["1", "2"], ["3", "4"]], sheet_title="X")
    sheet_names, chosen, headers, rows = csvimport.parse_xlsx_bytes(data)
    assert sheet_names == ["X"]
    assert chosen == "X"
    assert headers == ["A", "B"]
    assert rows == [{"A": "1", "B": "2"}, {"A": "3", "B": "4"}]


def test_parse_xlsx_bytes_unit_multi_sheet_without_choice():
    fresh()
    data = _multi_sheet_xlsx_bytes({
        "One": (["A"], [["1"]]),
        "Two": (["B"], [["2"]]),
    })
    sheet_names, chosen, headers, rows = csvimport.parse_xlsx_bytes(data)
    assert set(sheet_names) == {"One", "Two"}
    assert chosen is None
    assert headers is None
    assert rows is None
